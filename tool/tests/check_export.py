"""Integration checks against real exports and optional local reference data."""
import collections
import json
import math
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from common import APP,ARCHIVE,latest,read_json,write_json
from openpyxl import load_workbook


def main():
    if (ARCHIVE/'version.json').is_file():
        from archive_store import verified_current
        from version_compare import tags,load_snapshot,compare
        current=verified_current(ARCHIVE)
        for cid,p in current['operators'].items():
            assert cid==p['id'] and p['variants']
            assert '参数未提供' not in json.dumps(p,ensure_ascii=False),cid
        for mid,m in current['modules'].items():
            assert mid==m['id'] and m['operator_id'] in current['operators']
            assert '参数未提供' not in json.dumps(m,ensure_ascii=False),mid
            assert len({lev['level'] for lev in m['levels']})==len(m['levels'])
        if current['version']['official_version']=='78.2.0':
            clemnt=current['operators']['char_4231_clemnt']
            assert clemnt['variants'][0]['stats']['atk']==1030
            assert clemnt['variants'][-1]['stats']['atk']==1065
            assert '235%' in clemnt['skills'][2]['levels'][-1]['description']
        versions=tags(ARCHIVE)
        stored=load_snapshot(ARCHIVE,versions[-1])
        diff=compare(stored,current)
        assert not any(diff[k][kind] for k in ('operators','modules') for kind in ('added','removed','changed'))
        report={'official_version':current['version']['official_version'],'operators':len(current['operators']),
                'modules':len(current['modules']),'tags':versions,'working_tree_matches_tag':True,'unresolved_parameters':0}
        write_json(APP/'work/final-validation.json',report)
        print('Git archive and all operator/module checks passed:',json.dumps(report,ensure_ascii=False),flush=True)
        return
    run=latest();summary=read_json(run/'success.json');report={'workbooks':[],'reference_comparisons':[]}
    for spec in summary['workbooks']:
        path=run/'Excel'/spec['file'];wb=load_workbook(path,read_only=True,data_only=False)
        assert set(wb.sheetnames)==set(spec['sheets'])
        for name,expected in spec['sheets'].items():
            rows=wb[name].iter_rows();headers=[c.value for c in next(rows)];count=0
            for row in rows:
                assert all(c.data_type!='f' for c in row),f'Unexpected formula {path}/{name}'
                count+=1
            assert count==expected,(name,count,expected)
            assert len(headers)==len(set(headers)),f'Duplicate headers in {name}'
            report['workbooks'].append({'file':spec['file'],'sheet':name,'rows':count})
        wb.close()
    wb=load_workbook(run/'Excel/01_干员.xlsx',read_only=True,data_only=True)
    rows=wb['属性关键帧'].iter_rows(values_only=True);headers=next(rows);matches=[]
    for row in rows:
        item=dict(zip(headers,row))
        if item['角色ID']=='char_002_amiya' and item['精英阶段']==2 and item['等级']==80:matches.append(item)
    assert len(matches)==1
    sample=matches[0];assert sample['名称']=='阿米娅' and sample['生命']==1480 and sample['攻击']==612 and sample['防御']==121
    wb.close();report['sample_amiya']=sample
    schema=read_json(run/'schemas/schema.json')
    reference=APP.parent/'明日方舟国服数据_rel78/repo/zh_CN/gamedata/excel'
    for label,root in schema['roots'].items():
        ref=reference/(label+'.json')
        if label=='level' or not ref.exists():continue
        own=read_json(run/'json'/(label+'.json'));other=read_json(ref);checked=[0]
        def compare(typ,a,b,path):
            checked[0]+=1;n=schema['nodes'][typ];k=n['kind']
            if a is None or b is None:
                assert a==b,(path,a,b);return
            if k=='enum':
                if isinstance(b,str):b=n['values'].get(b,b)
                assert a==b,(path,a,b)
            elif k=='dict':
                keytype=schema['nodes'][n['key']]
                if keytype['kind']=='enum':b={str(keytype['values'].get(key,key)):value for key,value in b.items()}
                assert a.keys()==b.keys(),(path,'dictionary keys')
                for key in a:compare(n['value'],a[key],b[key],path+'/'+key)
            elif k=='vector':
                assert len(a)==len(b),(path,'vector length')
                for i,(x,y) in enumerate(zip(a,b)):compare(n['element'],x,y,path+'/'+str(i))
            elif k=='table':
                assert a.keys()==b.keys(),(path,'table fields',a.keys()^b.keys())
                for f in n['fields']:compare(f['type'],a[f['name']],b[f['name']],path+'/'+f['name'])
            elif isinstance(a,float):assert math.isclose(a,b,rel_tol=1e-6,abs_tol=1e-7),(path,a,b)
            else:assert a==b,(path,a,b)
        compare(root,own,other,label)
        report['reference_comparisons'].append({'table':label,'checked_nodes':checked[0],'equal':True})
        print(label,checked[0],'matched',flush=True)
    report['reference_note']='Optional development check only; extraction and normal verification do not use this reference repository.'
    write_json(APP/'work/final-validation.json',report)
    print('All workbook checks and available local reference comparisons passed',flush=True)


if __name__=='__main__':main()
