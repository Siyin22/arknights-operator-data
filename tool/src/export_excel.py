"""Human-readable spreadsheets; full fidelity data remains in JSON."""
import argparse
import json
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from common import read_json,write_json

PROF={1:'近卫',2:'狙击',4:'重装',8:'医疗',16:'辅助',32:'术师',64:'特种',128:'召唤物',256:'陷阱',512:'先锋'}
SKILL_TYPE={0:'被动',1:'手动触发',2:'自动触发'}
SP_TYPE={1:'自动回复',2:'攻击回复',4:'受击回复',8:'被动'}


def compact(x):
    return json.dumps(x,ensure_ascii=False,separators=(',',':')) if x is not None else None


def pick(d,*keys):return [d.get(k) for k in keys]


def defined(v):
    return v.get('m_value') if isinstance(v,dict) and v.get('m_defined') else None


class Excel:
    def __init__(self,path):
        self.path=path;self.book=Workbook(write_only=True);self.counts={};self.shortened=0

    def sheet(self,name,headers,width=20):
        ws=self.book.create_sheet(name);ws.freeze_panes='A2'
        for i,h in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(i)].width=44 if any(k in h for k in ['描述','JSON','说明','路径']) else width
        cells=[]
        for h in headers:
            cell=WriteOnlyCell(ws,h);cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='23465A');cells.append(cell)
        ws.append(cells);self.counts[name]=0
        return ws

    def row(self,ws,values):
        cells=[]
        for value in values:
            if isinstance(value,(list,dict)):value=compact(value)
            if isinstance(value,str):
                value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',lambda m:'\\u%04x'%ord(m[0]),value)
                if len(value)>32767:
                    value=value[:32680]+' …【超出 Excel 单元格长度，完整内容见对应 JSON】';self.shortened+=1
                # Keep game text literal, including strings beginning with '='.
                cell=WriteOnlyCell(ws,value);cell.data_type='s';cells.append(cell)
            else:cells.append(value)
        ws.append(cells);self.counts[ws.title]+=1
        if self.counts[ws.title]>=1048575:raise ValueError('Excel 行数超限，请拆分导出。')

    def save(self):
        for ws in self.book.worksheets:
            ws.auto_filter.ref=f'A1:{get_column_letter(len(ws.column_dimensions))}{self.counts[ws.title]+1}'
        self.book.save(self.path)
        return {'file':self.path.name,'sheets':self.counts,'shortened_cells':self.shortened}


def export(run):
    source=run/'json';out=run/'Excel';out.mkdir(exist_ok=True)
    load=lambda n:read_json(source/(n+'.json'))
    chars=load('character_table');skills=load('skill_table');equips=load('uniequip_table')
    charname=lambda cid:(chars.get(cid) or {}).get('name',cid)
    summaries=[]
    book=Excel(out/'01_干员.xlsx')
    ws=book.sheet('属性关键帧',['角色ID','记录类别','名称','星级','职业','分支','不可获取标记','精英阶段','等级','生命','攻击','防御','法抗','费用','阻挡','攻击间隔','再部署时间','技能ID列表'])
    tal=book.sheet('天赋候选',['干员ID','名称','天赋序号','天赋名','精英阶段','等级','所需潜能索引','描述','参数JSON'])
    sk=book.sheet('技能关联',['干员ID','名称','技能序号','技能ID','精英阶段','等级','其他条件JSON'])
    for cid,c in chars.items():
        for phase,p in enumerate(c.get('phases') or []):
            for frame in p.get('attributesKeyFrames') or []:
                a=frame['data']
                kind={'char':'干员','token':'召唤物','trap':'装置'}.get(cid.split('_')[0],'其他')
                book.row(ws,[cid,kind,c['name'],c['rarity']+1 if kind=='干员' else None,PROF.get(c['profession'],c['profession']),c['subProfessionId'],c['isNotObtainable'],phase,frame['level'],*pick(a,'maxHp','atk','def','magicResistance','cost','blockCnt','baseAttackTime','respawnTime'),','.join(s['skillId'] for s in c.get('skills') or [] if s.get('skillId'))])
        for i,t in enumerate(c.get('talents') or [],1):
            for candidate in t.get('candidates') or []:
                cond=candidate.get('unlockCondition') or {}
                book.row(tal,[cid,c['name'],i,candidate.get('name'),cond.get('phase'),cond.get('level'),candidate.get('requiredPotentialRank'),candidate.get('description'),candidate.get('blackboard')])
        for i,s in enumerate(c.get('skills') or [],1):
            cond=s.get('unlockCond') or s.get('unlockCondition') or {}
            book.row(sk,[cid,c['name'],i,s.get('skillId'),cond.get('phase'),cond.get('level'),s])
    summaries.append(book.save());print('干员 Excel 已完成',flush=True)
    book=Excel(out/'02_技能.xlsx')
    ws=book.sheet('技能等级',['技能ID','等级序号','名称','范围ID','触发方式','回复方式','SP消耗','初始SP','充能数','持续时间','持续类型原值','描述','参数JSON'])
    params=book.sheet('技能参数',['技能ID','等级序号','参数名','数值','字符串值'])
    for sid,s in skills.items():
        for level,d in enumerate(s.get('levels') or [],1):
            sp=d.get('spData') or {}
            book.row(ws,[sid,level,d.get('name'),d.get('rangeId'),SKILL_TYPE.get(d.get('skillType'),d.get('skillType')),SP_TYPE.get(sp.get('spType'),sp.get('spType')),*pick(sp,'spCost','initSp','maxChargeTime'),d.get('duration'),d.get('durationType'),d.get('description'),d.get('blackboard')])
            for v in d.get('blackboard') or []:book.row(params,[sid,level,*pick(v,'key','value','valueStr')])
    summaries.append(book.save());print('技能 Excel 已完成',flush=True)
    book=Excel(out/'03_模组.xlsx')
    ws=book.sheet('模组信息',['模组ID','名称','干员ID','干员名','类型1','类型2','解锁精英阶段','解锁等级','描述','升级材料JSON'])
    stats=book.sheet('属性加成',['模组ID','干员名','模组等级','对象','参数名','数值','字符串值'])
    effects=book.sheet('天赋特性变更',['模组ID','干员名','模组等级','部件序号','作用对象原值','是否召唤物','有效游戏标签JSON','有效地图标签JSON','天赋变更JSON','特性变更JSON'])
    for eid,e in equips['equipDict'].items():
        book.row(ws,[eid,e.get('uniEquipName'),e.get('charId'),charname(e.get('charId')),*pick(e,'typeName1','typeName2','unlockEvolvePhase','unlockLevel','uniEquipDesc','itemCost')])
    for eid,e in load('battle_equip_table').items():
        cname=charname((equips['equipDict'].get(eid) or {}).get('charId'))
        for phase in e.get('phases') or []:
            level=phase['equipLevel']
            for v in phase.get('attributeBlackboard') or []:book.row(stats,[eid,cname,level,'干员',*pick(v,'key','value','valueStr')])
            for token,values in (phase.get('tokenAttributeBlackboard') or {}).items():
                for v in values:book.row(stats,[eid,cname,level,token,*pick(v,'key','value','valueStr')])
            for i,p in enumerate(phase.get('parts') or [],1):
                book.row(effects,[eid,cname,level,i,*pick(p,'target','isToken','validInGameTag','validInMapTag','addOrOverrideTalentDataBundle','overrideTraitDataBundle')])
    summaries.append(book.save());print('模组 Excel 已完成',flush=True)
    book=Excel(out/'04_敌人.xlsx')
    ws=book.sheet('分级原始定义',['敌人ID','等级索引','名称定义','生命定义','攻击定义','防御定义','法抗定义','移速定义','攻击间隔定义','重量定义','扣除生命定义','描述定义','天赋参数JSON','技能JSON'])
    enemy_entries=load('enemy_database')['enemies']
    for e in enemy_entries:
        for level in e['value']:
            d=level['enemyData'];a=d.get('attributes') or {}
            book.row(ws,[e['key'],level['level'],defined(d.get('name')),*[defined(a.get(k)) for k in ['maxHp','atk','def','magicResistance','moveSpeed','baseAttackTime','massLevel']],defined(d.get('lifePointReduce')),defined(d.get('description')),d.get('talentBlackboard'),d.get('skills')])
    summaries.append(book.save());print('敌人 Excel 已完成',flush=True)
    book=Excel(out/'05_关卡.xlsx')
    ws=book.sheet('关卡总表',['关卡ID','代号','名称','区域ID','难度原值','难度组原值','理智','推荐等级','是否纯剧情','描述','关卡文件ID','掉落显示JSON'])
    stages=load('stage_table')['stages']
    for sid,s in stages.items():book.row(ws,[sid,*pick(s,'code','name','zoneId','difficulty','diffGroup','apCost','dangerLevel','isStoryOnly','description','levelId','stageDropInfo')])
    index=book.sheet('战斗文件索引',['相对路径','关卡ID字段','地图ID','地图行数','地图列数','路线数','波次数','分支数','敌人引用JSON'])
    actions=book.sheet('主波次行动',['关卡文件','波次序号','片段序号','行动序号','波前延迟','片段前延迟','行动类型原值','对象ID','数量','行动前延迟','间隔','路线索引','随机组','权重','隐藏组'])
    levelpaths=sorted((source/'levels').rglob('*.json'))
    for i,p in enumerate(levelpaths,1):
        d=read_json(p);rel=p.relative_to(source).as_posix();matrix=(d.get('mapData') or {}).get('map') or []
        book.row(index,[rel,d.get('levelId'),d.get('mapId'),len(matrix),len(matrix[0]) if matrix else 0,len(d.get('routes') or []),len(d.get('waves') or []),len(d.get('branches') or {}),d.get('enemyDbRefs')])
        for wi,w in enumerate(d.get('waves') or [],1):
            for fi,f in enumerate(w.get('fragments') or [],1):
                for ai,a in enumerate(f.get('actions') or [],1):book.row(actions,[rel,wi,fi,ai,w.get('preDelay'),f.get('preDelay'),*pick(a,'actionType','key','count','preDelay','interval','routeIndex','randomSpawnGroupKey','weight','hiddenGroup')])
        if i%500==0:print(f'关卡 Excel {i}/{len(levelpaths)}',flush=True)
    summaries.append(book.save())
    counts={'character_records':len(chars),'operator_records':sum(k.startswith('char_') for k in chars),'obtainable_characters':sum(k.startswith('char_') and not c['isNotObtainable'] for k,c in chars.items()),'skills':len(skills),'module_records':len(equips['equipDict']),'enemy_records':len(enemy_entries),'stages':len(stages),'level_files':len(levelpaths)}
    summary={'counts':counts,'workbooks':summaries}
    write_json(run/'export-summary.json',summary)
    note='''明日方舟数据研究工具 · 导出说明

首先打开“更新报告.html”：对比上一不同版本，只展示新增干员与模组；可离线阅读、切换潜能、打印为 PDF。
新增干员采用最高精英阶段、满级、满信赖属性，技能默认显示最高等级，可展开七级与其他专精等级。
新增模组展示每一级的属性与特性/天赋效果。更新报告.md 是一潜文字版，update-report.json 是结构化增量。
没有更早成功版本时只建立基准，不把全部内容当作新增；同版本重复提取仍对比上一个不同版本。
请保留上一版本的成功目录，以便后续自动对比。已有条目的数值修改不计为新增。

Excel：便于筛选、比较和计算的摘要表。json：完整配置数据，供后续研究工具读取。
schemas：从本机客户端生成的 .fbs 与字段结构；无需手动编辑。
raw：从本机当前资源清单提取的原始数据；analysis：供 AI 适配使用的分析中间文件。

使用前请了解：
1. 属性表包含干员、召唤物和装置，可按“记录类别”筛选。它是各精英阶段的属性关键帧，未叠加信赖、潜能、模组和战斗增益；不是实战面板。不可获取标记是客户端字段，不代表当前卡池。
2. 技能等级按数据数组从 1 开始计数；参数保留原始值，百分比和特殊效果需要结合描述解释。
3. 敌人分级表只显示该等级明确提供的定义。空白表示未定义，不等于 0，也不擅自计算继承后的面板。
4. 主波次行动表保留调度参数，不是已经计算好的出怪时间轴。随机、分支、条件和运行时变化以完整关卡 JSON 为准。
5. 原始 JSON 的枚举保留数值；含义可在 schemas/schema.json 的枚举定义中查到。字典键统一为字符串。
6. 描述中的游戏富文本标签和参数占位符保留。浮点数可能显示二进制表示误差，Excel 仅调整显示即可。
7. 客户端掉落显示字段不是服务器精确掉率，不应用它推算精确掉落概率。
8. Excel 单元格过长时标注省略，完整内容始终保留在 JSON。
9. 不需要联网、登录账号或运行游戏。导出期间请等游戏更新完成，避免资源版本发生变化。

提取失败时不要使用失败目录中的中间数据；打开上次成功结果，或将失败目录中的诊断包.zip交给 AI。
'''
    (run/'先读我.txt').write_text(note,encoding='utf-8-sig')
    print('Excel 导出完成：'+json.dumps(counts,ensure_ascii=False),flush=True)
    return summary


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('run',type=Path);a=ap.parse_args();export(a.run)
